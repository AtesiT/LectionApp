#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Сборка отчёта о программе Motion Playground Web.

Скрипт берёт содержание из tools/report_content.py и собирает два файла:

  report/Otchet_Motion_Playground_Web.html — самодостаточный HTML со встроенными
                                             рисунками (base64), его можно открыть
                                             браузером или Word и вставить в документ;
  report/Otchet_Motion_Playground_Web.docx — тот же отчёт в формате Word
                                             (нужен python-docx, иначе пропускается).

Запуск:  python3 tools/make_report.py
Рисунки берутся из report/img/ (их собирает tools/make_report_images.py).
"""

import base64
import html
import os
import re
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from report_content import DOC, TITLE, SUBTITLE          # noqa: E402

IMG_DIR = os.path.join(ROOT, 'report', 'img')
OUT_DIR = os.path.join(ROOT, 'report')
HTML_PATH = os.path.join(OUT_DIR, 'Otchet_Motion_Playground_Web.html')
DOCX_PATH = os.path.join(OUT_DIR, 'Otchet_Motion_Playground_Web.docx')

PAGE_WIDTH_CM = 17.0          # ширина текстового поля A4 при полях 2 см
IMG_WIDTH_CM = 16.2

# ---------------------------------------------------------------------------
# Общее: разметка внутри строк (**жирный**, `код`)
# ---------------------------------------------------------------------------

def inline(text):
    s = html.escape(str(text), quote=False)
    s = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', s, flags=re.S)
    s = re.sub(r'`([^`]+)`',
               r'<span style="font-family:Consolas,\'Courier New\',monospace;font-size:11.5pt;">\1</span>',
               s)
    return s.replace('\n', '<br>')


def plain(text):
    """Убирает разметку **жирный** и `код` — для Word и обычного текста."""
    s = re.sub(r'\*\*(.+?)\*\*', r'\1', str(text), flags=re.S)
    return re.sub(r'`([^`]+)`', r'\1', s)


def runs(text):
    """Разбивает строку на куски [(текст, жирный, моно)]."""
    out = []
    pos = 0
    for m in re.finditer(r'\*\*(.+?)\*\*|`([^`]+)`', str(text), flags=re.S):
        if m.start() > pos:
            out.append((text[pos:m.start()], False, False))
        if m.group(1) is not None:
            out.append((m.group(1), True, False))
        else:
            out.append((m.group(2), False, True))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], False, False))
    return out or [(str(text), False, False)]


def figure_files():
    used = [b[1] for b in DOC if b[0] == 'fig']
    return used


# ---------------------------------------------------------------------------
# HTML
# ---------------------------------------------------------------------------

CSS = """
@page { size: A4; margin: 2cm; }
body { font-family:'Times New Roman',Times,serif; font-size:14pt; line-height:1.5;
       color:#111111; margin:0; }
h1 { font-size:22pt; font-weight:bold; text-align:center; margin:0 0 6pt; }
h2 { font-size:16pt; font-weight:bold; margin:16pt 0 6pt; text-align:left; }
h3 { font-size:14pt; font-weight:bold; margin:12pt 0 4pt; }
p  { margin:0 0 6pt; text-align:justify; text-indent:1.25cm; }
p.noindent { text-indent:0; }
div.title-sub { text-align:center; font-size:15pt; margin:0 0 18pt; }
div.meta { text-align:center; font-size:12pt; margin:0 0 4pt; }
ul, ol { margin:0 0 8pt 1.25cm; padding:0; }
li { margin:0 0 3pt; }
table { width:100%; border-collapse:collapse; margin:6pt 0 12pt; font-size:11.5pt; }
td, th { border:1px solid #333333; padding:3pt 6pt; vertical-align:top; }
th { background:#eef2f7; font-weight:bold; text-align:left; }
img { display:block; margin:6pt auto 2pt; height:auto; }
p.caption { text-align:center; font-style:italic; font-size:11pt; text-indent:0;
            margin:0 0 12pt; }
p.code { font-family:Consolas,'Courier New',monospace; font-size:11pt; text-indent:0;
         background:#f4f6f8; padding:6pt 8pt; border:1px solid #cccccc; }
div.pb { page-break-before:always; }
"""


def b64(name):
    with open(os.path.join(IMG_DIR, name), 'rb') as fh:
        return base64.b64encode(fh.read()).decode('ascii')


def table_html(header, rows, widths=None):
    out = ['<table>']
    if widths:
        out.append('<colgroup>')
        for w in widths:
            out.append(f'<col style="width:{w};">')
        out.append('</colgroup>')
    if header:
        out.append('<tr>' + ''.join(f'<th>{inline(c)}</th>' for c in header) + '</tr>')
    for row in rows:
        out.append('<tr>' + ''.join(f'<td>{inline(c)}</td>' for c in row) + '</tr>')
    out.append('</table>')
    return '\n'.join(out)


def block_html(b, cache):
    kind = b[0]
    if kind == 'h1':
        return f'<h1>{inline(b[1])}</h1>'
    if kind == 'h2':
        return f'<h2>{inline(b[1])}</h2>'
    if kind == 'h3':
        return f'<h3>{inline(b[1])}</h3>'
    if kind == 'p':
        cls = ' class="noindent"' if len(b) > 2 and b[2] == 'noindent' else ''
        return f'<p{cls}>{inline(b[1])}</p>'
    if kind == 'center':
        return f'<p class="noindent" style="text-align:center">{inline(b[1])}</p>'
    if kind == 'code':
        return f'<p class="code">{inline(b[1])}</p>'
    if kind == 'ul':
        items = '\n'.join(f'<li>{inline(i)}</li>' for i in b[1])
        return f'<ul>{items}</ul>'
    if kind == 'ol':
        items = '\n'.join(f'<li>{inline(i)}</li>' for i in b[1])
        return f'<ol>{items}</ol>'
    if kind == 'table':
        header = b[1]
        rows = b[2]
        widths = b[3] if len(b) > 3 else None
        return table_html(header, rows, widths)
    if kind == 'fig':
        name, caption = b[1], b[2]
        if name not in cache:
            cache[name] = b64(name)
        src = 'data:image/png;base64,' + cache[name]
        return (f'<p class="noindent" style="text-align:center">'
                f'<img src="{src}" style="width:{IMG_WIDTH_CM}cm;max-width:100%;height:auto;">'
                f'</p>\n<p class="caption">{inline(caption)}</p>')
    if kind == 'pb':
        return '<div class="pb"></div>'
    if kind == 'fields':
        out = []
        for label in b[1]:
            out.append(f'<p class="noindent" style="margin:0 0 14pt">{inline(label)}</p>')
        return '\n'.join(out)
    raise ValueError('неизвестный блок: ' + str(kind))


def build_html():
    cache = {}
    parts = ['<!DOCTYPE html>', '<html lang="ru">', '<head>',
             '<meta charset="utf-8">',
             f'<title>{html.escape(TITLE)}</title>',
             f'<style>{CSS}</style>', '</head>', '<body>']
    for b in DOC:
        parts.append(block_html(b, cache))
    parts.append('</body></html>')
    text = '\n'.join(parts)
    os.makedirs(OUT_DIR, exist_ok=True)
    # BOM нужен, чтобы Word правильно распознал кириллицу при открытии файла
    with open(HTML_PATH, 'w', encoding='utf-8-sig') as fh:
        fh.write(text)
    return len(text)


# ---------------------------------------------------------------------------
# DOCX (python-docx, если установлен)
# ---------------------------------------------------------------------------

def build_docx():
    try:
        from docx import Document
        from docx.shared import Pt, Cm, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
    except ImportError:
        print('  python-docx не установлен — DOCX пропущен (HTML готов)')
        return False

    doc = Document()

    # --- параметры страницы ---
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)

    # --- стиль «Обычный» ---
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(14)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)

    def tune(style_name, size, bold=True, center=False, before=12, after=6):
        st = doc.styles[style_name]
        st.font.name = 'Times New Roman'
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
        st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.line_spacing = 1.15
        if center:
            st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    tune('Heading 1', 20, center=True, before=0, after=10)
    tune('Heading 2', 16, before=14, after=6)
    tune('Heading 3', 14, before=10, after=4)

    for style_name in ('List Bullet', 'List Number'):
        try:
            st = doc.styles[style_name]
            st.font.name = 'Times New Roman'
            st.font.size = Pt(14)
            st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
        except KeyError:
            pass

    def add_runs(par, text, size=None, italic=False):
        for chunk, bold, mono in runs(text):
            r = par.add_run(chunk)
            r.font.name = 'Times New Roman' if not mono else 'Consolas'
            if size:
                r.font.size = Pt(size)
            r.bold = bold
            r.italic = italic
            if mono:
                r.font.size = Pt(11.5)
        return par

    def add_page_number(par):
        run = par.add_run()
        f1 = OxmlElement('w:fldChar')
        f1.set(qn('w:fldCharType'), 'begin')
        it = OxmlElement('w:instrText')
        it.set(qn('xml:space'), 'preserve')
        it.text = 'PAGE'
        f2 = OxmlElement('w:fldChar')
        f2.set(qn('w:fldCharType'), 'end')
        run._r.append(f1)
        run._r.append(it)
        run._r.append(f2)

    # --- колонтитул с номером страницы ---
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run()
    fr.font.name = 'Times New Roman'
    fr.font.size = Pt(11)
    add_page_number(footer)

    content_width = int(PAGE_WIDTH_CM * 360000)   # в EMU

    for b in DOC:
        kind = b[0]
        if kind == 'h1':
            doc.add_paragraph(plain(b[1]), style='Heading 1')
        elif kind == 'h2':
            doc.add_paragraph(plain(b[1]), style='Heading 2')
        elif kind == 'h3':
            doc.add_paragraph(plain(b[1]), style='Heading 3')
        elif kind == 'p':
            par = doc.add_paragraph()
            par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            par.paragraph_format.first_line_indent = Cm(1.25)
            add_runs(par, b[1])
        elif kind == 'center':
            par = doc.add_paragraph()
            par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(par, b[1])
        elif kind == 'code':
            par = doc.add_paragraph()
            add_runs(par, b[1])
            par.paragraph_format.left_indent = Cm(0.6)
        elif kind in ('ul', 'ol'):
            for item in b[1]:
                par = doc.add_paragraph(style='List Bullet' if kind == 'ul' else 'List Number')
                par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                add_runs(par, item)
        elif kind == 'table':
            header, rows = b[1], b[2]
            widths = b[3] if len(b) > 3 else None
            ncol = len(header) if header else len(rows[0])
            table = doc.add_table(rows=0, cols=ncol)
            table.style = 'Table Grid'
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            if header:
                cells = table.add_row().cells
                for i, c in enumerate(header):
                    cells[i].text = ''
                    par = cells[i].paragraphs[0]
                    par.paragraph_format.space_after = Pt(2)
                    add_runs(par, c, size=11.5)
                    for r in par.runs:
                        r.bold = True
            for row in rows:
                cells = table.add_row().cells
                for i, c in enumerate(row):
                    cells[i].text = ''
                    par = cells[i].paragraphs[0]
                    par.paragraph_format.space_after = Pt(2)
                    add_runs(par, c, size=11.5)
            if widths:
                total = sum(widths)
                for row in table.rows:
                    for i, w in enumerate(widths):
                        if i < len(row.cells):
                            row.cells[i].width = int(content_width * w / total)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
        elif kind == 'fig':
            name, caption = b[1], b[2]
            par = doc.add_paragraph()
            par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.first_line_indent = Cm(0)
            run = par.add_run()
            run.add_picture(os.path.join(IMG_DIR, name), width=Cm(IMG_WIDTH_CM))
            cap = doc.add_paragraph()
            cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.keep_with_next = True
            add_runs(cap, caption, size=11, italic=True)
        elif kind == 'pb':
            doc.add_page_break()
        elif kind == 'fields':
            for label in b[1]:
                par = doc.add_paragraph()
                par.paragraph_format.space_after = Pt(12)
                add_runs(par, label)

    doc.save(DOCX_PATH)
    return True


# ---------------------------------------------------------------------------

def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    missing = [n for n in figure_files() if not os.path.exists(os.path.join(IMG_DIR, n))]
    if missing:
        print('Нет рисунков:', ', '.join(missing))
        print('Сначала запустите: python3 tools/make_report_images.py')
        return 1
    size = build_html()
    print(f'  {os.path.relpath(HTML_PATH, ROOT)}  ({size // 1024} КБ, рисунки встроены)')
    if build_docx():
        print(f'  {os.path.relpath(DOCX_PATH, ROOT)}  '
              f'({os.path.getsize(DOCX_PATH) // 1024} КБ)')
    print('Отчёт собран.')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
