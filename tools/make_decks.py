#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Пять разных презентаций, по 30 слайдов в каждой.

Презентации отличаются не только оформлением, но и содержанием: у каждой своя
аудитория, своя цель и собственный набор слайдов. Текст каждой презентации
лежит в отдельном модуле:

    deck1_zashchita.py   — защита проекта перед комиссией   (классическая)
    deck2_znakomstvo.py  — знакомство с программой, 7-9 класс (современная)
    deck3_uroki.py       — двенадцать уроков для учителя      (академическая)
    deck4_tehnika.py     — технический разбор для кружка      (тёмная)
    deck5_pitch.py       — короткий показ на выставке         (минимализм)

Каждый модуль отдаёт словарь DECK с полями key, file, theme, name, desc,
title, sub, meta, note и slides. Оформление задаётся словарём THEMES: шрифт,
цвета, decoration заголовка и вид номера слайда.

Результат — папка report/prezentacii: пять пар .html + .pptx и index.html.

Запуск:  python3 tools/make_decks.py
"""

import base64
import html
import os
import re
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from deck1_zashchita import DECK as DECK_1  # noqa: E402
from deck2_znakomstvo import DECK as DECK_2  # noqa: E402
from deck3_uroki import DECK as DECK_3  # noqa: E402
from deck4_tehnika import DECK as DECK_4  # noqa: E402
from deck5_pitch import DECK as DECK_5  # noqa: E402

DECKS = [DECK_1, DECK_2, DECK_3, DECK_4, DECK_5]

IMG_DIR = os.path.join(ROOT, 'report', 'img')
OUT_DIR = os.path.join(ROOT, 'report', 'prezentacii')

# ---------------------------------------------------------------------------
# Оформления
# ---------------------------------------------------------------------------

THEMES = [
    dict(key='klassicheskaya', name='Классическая', file='Prezentaciya_1_klassicheskaya',
         desc='Белый фон, Times New Roman, тонкая полоса сверху и аккуратная линейка под '
              'заголовком. Спокойное, «экзаменационное» оформление.',
         font='Times New Roman', title_size=30, body_size=17, body_two=17,
         bg='FFFFFF', title_c='0F172A', body_c='1F2937', accent='1D4ED8',
         muted='6B7280', marker='• ', decor='bar', number='footer',
         title_bg='0B1020', title_c2='FFFFFF', sub_c='93C5FD', meta_c='C7D2DE'),

    dict(key='temnaya', name='Тёмная', file='Prezentaciya_2_temnaya',
         desc='Тёмно-графитовый фон, Verdana и бирюзовый акцент: хорошо смотрится на '
              'проекторе в затемнённом классе.',
         font='Verdana', title_size=29, body_size=16, body_two=16,
         bg='0F172A', title_c='FFFFFF', body_c='CBD5E1', accent='22D3EE',
         muted='64748B', marker='— ', decor='leftbar', number='footer',
         title_bg='0F172A', title_c2='FFFFFF', sub_c='67E8F9', meta_c='94A3B8'),

    dict(key='akademicheskaya', name='Академическая', file='Prezentaciya_3_akademicheskaya',
         desc='Кремовый фон, Georgia, бордовый заголовок и тонкая рамка по периметру слайда.',
         font='Georgia', title_size=29, body_size=16, body_two=16,
         bg='FDF8F0', title_c='7F1D1D', body_c='3F2D1E', accent='B45309',
         muted='8A7355', marker='— ', decor='frame', number='footer',
         title_bg='FDF8F0', title_c2='7F1D1D', sub_c='92400E', meta_c='6B5844'),

    dict(key='sovremennaya', name='Современная', file='Prezentaciya_4_sovremennaya',
         desc='Широкая цветная полоса с заголовком, шрифт Calibri, фиолетовый маркер и '
              'номер слайда в цветной плашке.',
         font='Calibri', title_size=30, body_size=18, body_two=18,
         bg='FFFFFF', title_c='FFFFFF', body_c='1F2937', accent='7C3AED',
         muted='6B7280', marker='▸ ', decor='band', number='badge',
         band_from='2563EB', band_to='7C3AED',
         title_bg=None, title_c2='FFFFFF', sub_c='C4B5FD', meta_c='E2E8F0',
         title_from='1E3A8A', title_to='7C3AED'),

    dict(key='minimalizm', name='Минимализм', file='Prezentaciya_5_minimalizm',
         desc='Много свободного места, Verdana, зелёная точка-акцент и крупный номер слайда '
              'в углу. Пункты без маркеров.',
         font='Verdana', title_size=27, body_size=16, body_two=16,
         bg='FFFFFF', title_c='111111', body_c='374151', accent='16A34A',
         muted='9CA3AF', marker='', decor='minimal', number='biggray',
         title_bg='FFFFFF', title_c2='111111', sub_c='4B5563', meta_c='6B7280'),
]

THEME_BY_KEY = {t['key']: t for t in THEMES}


def deck_slides(deck):
    """Титульный слайд плюс слайды самой презентации."""
    return ([dict(kind='title', title=deck['title'], sub=deck['sub'],
                  meta=deck['meta'], note=deck['note'])]
            + [dict(kind='slide', **s) for s in deck['slides']])

# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------

def plain(text):
    return re.sub(r'\*\*(.+?)\*\*', r'\1', str(text))


def inline(text):
    return re.sub(r'\*\*(.+?)\*\*', r'<b>\\1</b>', html.escape(str(text), quote=False))


def png_size(path):
    with open(path, 'rb') as fh:
        head = fh.read(26)
    if head[:8] != b'\x89PNG\r\n\x1a\n':
        return 1240, 700
    return int.from_bytes(head[16:20], 'big'), int.from_bytes(head[20:24], 'big')


def b64(name, cache):
    if name not in cache:
        with open(os.path.join(IMG_DIR, name), 'rb') as fh:
            cache[name] = base64.b64encode(fh.read()).decode('ascii')
    return cache[name]


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------

def build_pptx(theme, path, slides, deck):
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt, Emu
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
        from pptx.enum.shapes import MSO_SHAPE
    except ImportError:
        print('  python-pptx не установлен — PPTX пропущен')
        return False

    def rgb(hex_str):
        h = hex_str.lstrip('#')
        return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

    A = rgb(theme['accent'])
    BG = rgb(theme['bg'])
    TITLE_C = rgb(theme['title_c'])
    BODY_C = rgb(theme['body_c'])
    MUTED_C = rgb(theme['muted'])
    FONT = theme['font']

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def rect(slide, x, y, w, h, color):
        sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
        sh.fill.solid()
        sh.fill.fore_color.rgb = color
        sh.line.fill.background()
        sh.shadow.inherit = False
        return sh

    def gradient_rect(slide, x, y, w, h, c1, c2):
        sh = rect(slide, x, y, w, h, rgb(c1))
        try:
            sh.fill.gradient()
            stops = sh.fill.gradient_stops
            stops[0].color.rgb = rgb(c1)
            stops[0].position = 0.0
            stops[1].color.rgb = rgb(c2)
            stops[1].position = 1.0
        except Exception:
            sh.fill.solid()
            sh.fill.fore_color.rgb = rgb(c1)
        return sh

    def textbox(slide, x, y, w, h, align=PP_ALIGN.LEFT):
        box = slide.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = 0
        tf.margin_right = 0
        tf.margin_top = 0
        tf.margin_bottom = 0
        if tf.paragraphs:
            tf.paragraphs[0].alignment = align
        return box

    def put(box, lines, size, color, bold=False, space_after=8, spacing=1.0,
            bullet='• ', align=PP_ALIGN.LEFT, font=None):
        tf = box.text_frame
        tf.clear()
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if (i == 0 and tf.paragraphs) else tf.add_paragraph()
            p.text = (bullet + line) if bullet else line
            p.space_after = Pt(space_after)
            p.line_spacing = spacing
            p.alignment = align
            for run in p.runs:
                run.font.size = Pt(size)
                run.font.bold = bold
                run.font.color.rgb = color
                run.font.name = font or FONT

    # --- геометрия --------------------------------------------------------
    band_h = Inches(1.45) if theme['decor'] == 'band' else 0
    content_top = Inches(1.95) if theme['decor'] == 'band' else Inches(1.80)
    left = Inches(0.85) if theme['decor'] == 'leftbar' else Inches(0.75)
    content_w = Inches(11.6) if theme['decor'] == 'leftbar' else Inches(11.8)
    content_h = Inches(4.40) if theme['decor'] == 'band' else Inches(4.55)

    for idx, s in enumerate(slides, start=1):
        slide = prs.slides.add_slide(blank)

        # фон
        if theme['decor'] in ('band', 'bar', 'leftbar', 'minimal'):
            rect(slide, 0, 0, prs.slide_width, prs.slide_height, BG)
        else:  # frame
            rect(slide, 0, 0, prs.slide_width, prs.slide_height, BG)
            frame = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.3), Inches(0.3),
                                           Inches(12.733), Inches(6.9))
            frame.fill.background()
            frame.line.color.rgb = rgb('D6C3A5')
            frame.line.width = Pt(1.25)
            frame.shadow.inherit = False

        # --- титульный слайд ---------------------------------------------
        if s['kind'] == 'title':
            if theme.get('title_from'):
                gradient_rect(slide, 0, 0, prs.slide_width, prs.slide_height,
                              theme['title_from'], theme['title_to'])
            else:
                rect(slide, 0, 0, prs.slide_width, prs.slide_height, rgb(theme['title_bg']))
            if theme['decor'] in ('bar', 'leftbar'):
                rect(slide, 0, 0, Inches(0.26), prs.slide_height, A)
            if theme['decor'] == 'frame':
                fr = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(0.55),
                                            Inches(12.233), Inches(6.4))
                fr.fill.background()
                fr.line.color.rgb = rgb('D6C3A5')
                fr.line.width = Pt(1.25)
                fr.shadow.inherit = False
            if theme['decor'] == 'minimal':
                rect(slide, Inches(1.0), Inches(1.75), Inches(0.18), Inches(0.18), A)
            box = textbox(slide, Inches(1.1), Inches(2.05), Inches(11.1), Inches(1.5))
            put(box, [s['title']], 44, rgb(theme['title_c2']), bold=True, bullet='',
                space_after=0)
            box = textbox(slide, Inches(1.1), Inches(3.55), Inches(10.4), Inches(1.3))
            put(box, [s['sub']], 22, rgb(theme['sub_c']), bullet='', space_after=0, spacing=1.2)
            box = textbox(slide, Inches(1.1), Inches(5.15), Inches(10.4), Inches(1.6))
            put(box, [plain(x) for x in s['meta']], 16, rgb(theme['meta_c']), bullet='',
                space_after=8, spacing=1.1)
            continue

        # --- обычный слайд ------------------------------------------------
        if theme['decor'] == 'band':
            gradient_rect(slide, 0, 0, prs.slide_width, band_h,
                          theme['band_from'], theme['band_to'])
            box = textbox(slide, Inches(0.8), Inches(0.30), Inches(11.7), Inches(0.85),
                          PP_ALIGN.LEFT)
            put(box, [s['title']], theme['title_size'], TITLE_C, bold=True, bullet='',
                space_after=0)
        elif theme['decor'] == 'bar':
            rect(slide, 0, 0, prs.slide_width, Inches(0.12), A)
            box = textbox(slide, left, Inches(0.45), content_w, Inches(0.9))
            put(box, [s['title']], theme['title_size'], TITLE_C, bold=True, bullet='',
                space_after=0)
            rect(slide, left, Inches(1.42), Inches(1.9), Inches(0.06), A)
        elif theme['decor'] == 'leftbar':
            rect(slide, 0, 0, Inches(0.22), prs.slide_height, A)
            box = textbox(slide, left, Inches(0.45), content_w, Inches(0.9))
            put(box, [s['title']], theme['title_size'], TITLE_C, bold=True, bullet='',
                space_after=0)
            rect(slide, left, Inches(1.42), Inches(2.2), Inches(0.05), A)
        elif theme['decor'] == 'frame':
            box = textbox(slide, Inches(0.85), Inches(0.55), Inches(11.6), Inches(0.9))
            put(box, [s['title']], theme['title_size'], TITLE_C, bold=True, bullet='',
                space_after=0)
            rect(slide, Inches(0.85), Inches(1.52), Inches(2.2), Inches(0.04), A)
            rect(slide, Inches(0.85), Inches(1.60), Inches(2.2), Inches(0.04), A)
        else:  # minimal
            rect(slide, Inches(0.85), Inches(0.62), Inches(0.16), Inches(0.16), A)
            box = textbox(slide, Inches(0.85), Inches(1.00), Inches(10.4), Inches(0.9))
            put(box, [s['title']], theme['title_size'], TITLE_C, bold=True, bullet='',
                space_after=0)
            num = textbox(slide, Inches(11.4), Inches(0.45), Inches(1.5), Inches(1.4),
                          PP_ALIGN.RIGHT)
            put(num, [str(idx)], 54, rgb('E2E8F0'), bold=True, bullet='', space_after=0,
                align=PP_ALIGN.RIGHT)

        bullets = [plain(b) for b in s.get('bullets', [])]
        marker = theme['marker']
        if s.get('img'):
            col_w = Inches(6.25) if theme['decor'] != 'leftbar' else Inches(6.15)
            box = textbox(slide, left, content_top, col_w, content_h)
            size = theme['body_size'] - 1 if len(bullets) >= 6 else theme['body_size']
            put(box, bullets, size, BODY_C, spacing=1.05, space_after=10, bullet=marker)

            ipath = os.path.join(IMG_DIR, s['img'])
            iw, ih = png_size(ipath)
            bw, bh = Inches(5.3), content_h
            pleft = Inches(7.4)
            scale = min(bw / iw, bh / ih)
            w, h = int(iw * scale), int(ih * scale)
            pic = slide.shapes.add_picture(ipath, pleft + int((bw - w) / 2),
                                           content_top + int((bh - h) / 2), w, h)
            if theme['decor'] in ('leftbar', 'frame', 'minimal'):
                pic.line.color.rgb = rgb('CBD5E1') if theme['decor'] == 'minimal' else A
                pic.line.width = Pt(0.75)
        else:
            if len(bullets) > 4:
                half = (len(bullets) + 1) // 2
                box = textbox(slide, left, content_top, Inches(5.6), content_h)
                put(box, bullets[:half], theme['body_two'], BODY_C, spacing=1.05,
                    space_after=12, bullet=marker)
                box = textbox(slide, Inches(7.0), content_top, Inches(5.5), content_h)
                put(box, bullets[half:], theme['body_two'], BODY_C, spacing=1.05,
                    space_after=12, bullet=marker)
            else:
                box = textbox(slide, left, content_top, content_w, content_h)
                put(box, bullets, theme['body_size'] + 1, BODY_C, spacing=1.15,
                    space_after=16, bullet=marker)

        # --- подвал --------------------------------------------------------
        if theme['number'] == 'footer':
            line_c = rgb('1E293B') if theme['decor'] == 'leftbar' else rgb('CBD5E1')
            rect(slide, left, Inches(6.95), content_w, Emu(9525), line_c)
            box = textbox(slide, left, Inches(7.05), Inches(8.0), Inches(0.35))
            put(box, [deck['title']], 11, MUTED_C, bullet='', space_after=0)
            box = textbox(slide, Inches(11.3), Inches(7.05), Inches(1.3), Inches(0.35),
                          PP_ALIGN.RIGHT)
            put(box, [str(idx)], 11, MUTED_C, bullet='', space_after=0, align=PP_ALIGN.RIGHT)
        elif theme['number'] == 'badge':
            badge = rect(slide, Inches(12.15), Inches(6.75), Inches(0.85), Inches(0.5), A)
            tf = badge.text_frame
            tf.word_wrap = False
            tf.margin_left = 0
            tf.margin_right = 0
            tf.margin_top = 0
            tf.margin_bottom = 0
            put(badge, [str(idx)], 14, RGBColor(0xFF, 0xFF, 0xFF), bold=True, bullet='',
                space_after=0, align=PP_ALIGN.CENTER, font=FONT)
            box = textbox(slide, left, Inches(7.02), Inches(8.0), Inches(0.35))
            put(box, [deck['title']], 11, MUTED_C, bullet='', space_after=0)

        if s.get('note'):
            slide.notes_slide.notes_text_frame.text = plain(s['note'])

    prs.save(path)
    return True


# ---------------------------------------------------------------------------
# HTML
# ---------------------------------------------------------------------------

BASE_CSS = """
* { box-sizing: border-box; }
html, body { margin:0; padding:0; height:100%; background:__BG__; overflow:hidden;
             font-family:__FONT__, serif; }
#stage { position:fixed; inset:0; display:flex; align-items:center; justify-content:center; }
#scaler { width:1280px; height:720px; transform-origin:center center; position:relative; }
.slide { position:absolute; inset:0; width:1280px; height:720px; background:__BG__;
         display:none; overflow:hidden; box-shadow:0 10px 40px rgba(0,0,0,.45); }
.slide.on { display:block; }
.slide h1 { position:absolute; margin:0; font-weight:bold; color:__TITLE__; }
.slide ul { position:absolute; margin:0; padding:0; list-style:none; }
.slide li { color:__BODY__; position:relative; }
.slide li::before { color:__ACCENT__; position:absolute; left:-26px; font-weight:bold; }
.slide .pic { position:absolute; display:flex; align-items:center; justify-content:center; }
.slide .pic img { max-width:100%; max-height:100%; height:auto; }
.slide .foot { position:absolute; height:1px; background:__RULE__; }
.slide .footname { position:absolute; font-size:15px; color:__MUTED__; }
.slide .num { position:absolute; font-size:15px; color:__MUTED__; }
.slide.title h1 { font-size:62px; }
.slide.title .sub { position:absolute; font-size:28px; color:__SUB__; }
.slide.title .meta { position:absolute; font-size:20px; line-height:1.7; color:__META__; }
#ui { position:fixed; right:18px; bottom:16px; display:flex; gap:8px; align-items:center;
      z-index:10; }
#ui button { font:16px/1 __FONT__, serif; padding:8px 14px; border:1px solid #334155;
             background:#111827; color:#e5e7eb; border-radius:6px; cursor:pointer; }
#ui button:hover { background:#1f2937; }
#count { color:#e5e7eb; font:16px/1 __FONT__, serif; padding:0 6px; }
#hint { position:fixed; left:18px; bottom:16px; color:#64748b; font:14px/1 __FONT__, serif;
        z-index:10; }
@media print {
  html, body { overflow:visible; background:#ffffff; }
  #stage { position:static; display:block; }
  #scaler { transform:none !important; width:auto; height:auto; }
  .slide { display:block !important; position:relative; page-break-after:always;
           box-shadow:none; }
  #ui, #hint { display:none; }
}
@page { size: 1280px 720px; margin:0; }
"""

# позиции для декора каждого оформления
LAYOUT = {
    'bar': """
.slide .topbar { position:absolute; left:0; top:0; width:1280px; height:12px;
                 background:__ACCENT__; }
.slide h1 { left:56px; top:38px; width:1170px; font-size:40px; line-height:1.15; }
.slide .rule { position:absolute; left:56px; top:118px; width:130px; height:5px;
               background:__ACCENT__; }
.slide ul { left:56px; top:168px; width:590px; }
.slide ul.wide { width:1170px; }
.slide ul.two { width:560px; }
.slide ul.right { left:650px; width:575px; }
.slide li { font-size:23px; line-height:1.32; margin:0 0 15px 26px; }
.slide li::before { content:"\\2022"; }
.slide .pic { left:660px; top:168px; width:575px; height:470px; }
.slide .foot { left:56px; bottom:44px; width:1170px; }
.slide .footname { left:56px; bottom:22px; }
.slide .num { right:56px; bottom:22px; }
.slide.title { background:__TITLEBG__; }
.slide.title .stripe { position:absolute; left:0; top:0; width:26px; height:720px;
                       background:__ACCENT__; }
.slide.title h1 { left:96px; top:200px; width:1100px; color:__TITLE2__; }
.slide.title .sub { left:96px; top:330px; width:1000px; }
.slide.title .meta { left:96px; top:470px; width:1000px; }
""",
    'leftbar': """
.slide .leftbar { position:absolute; left:0; top:0; width:22px; height:720px;
                  background:__ACCENT__; }
.slide h1 { left:82px; top:38px; width:1150px; font-size:38px; line-height:1.15; }
.slide .rule { position:absolute; left:82px; top:118px; width:150px; height:5px;
               background:__ACCENT__; }
.slide ul { left:82px; top:168px; width:585px; }
.slide ul.wide { width:1150px; }
.slide ul.two { width:555px; }
.slide ul.right { left:672px; width:560px; }
.slide li { font-size:22px; line-height:1.36; margin:0 0 14px 26px; }
.slide li::before { content:"\\2014"; }
.slide .pic { left:672px; top:168px; width:560px; height:470px; }
.slide .pic img { border:1px solid __ACCENT__; }
.slide .foot { left:82px; bottom:44px; width:1150px; background:#1e293b; }
.slide .footname { left:82px; bottom:22px; }
.slide .num { right:56px; bottom:22px; }
.slide.title { background:__TITLEBG__; }
.slide.title .stripe { position:absolute; left:0; top:0; width:26px; height:720px;
                       background:__ACCENT__; }
.slide.title h1 { left:96px; top:200px; width:1100px; color:__TITLE2__; }
.slide.title .sub { left:96px; top:330px; width:1000px; }
.slide.title .meta { left:96px; top:470px; width:1000px; }
""",
    'frame': """
.slide .frame { position:absolute; left:28px; top:28px; width:1224px; height:664px;
                border:1.5px solid #d6c3a5; }
.slide h1 { left:82px; top:52px; width:1120px; font-size:38px; line-height:1.16; }
.slide .rule { position:absolute; left:82px; top:132px; width:150px; height:3px;
               background:__ACCENT__; }
.slide .rule2 { position:absolute; left:82px; top:139px; width:150px; height:3px;
                background:__ACCENT__; }
.slide ul { left:82px; top:180px; width:585px; }
.slide ul.wide { width:1120px; }
.slide ul.two { width:555px; }
.slide ul.right { left:672px; width:530px; }
.slide li { font-size:22px; line-height:1.36; margin:0 0 14px 26px; }
.slide li::before { content:"\\2014"; }
.slide .pic { left:672px; top:180px; width:530px; height:450px; }
.slide .pic img { border:1px solid __ACCENT__; }
.slide .foot { left:82px; bottom:44px; width:1120px; }
.slide .footname { left:82px; bottom:22px; }
.slide .num { right:82px; bottom:22px; }
.slide.title { background:__TITLEBG__; }
.slide.title .frame { border:1.5px solid #d6c3a5; }
.slide.title h1 { left:110px; top:200px; width:1060px; color:__TITLE2__; }
.slide.title .sub { left:110px; top:330px; width:1000px; }
.slide.title .meta { left:110px; top:470px; width:1000px; }
""",
    'band': """
.slide .band { position:absolute; left:0; top:0; width:1280px; height:105px;
               background:linear-gradient(90deg, __BAND1__, __BAND2__); }
.slide h1 { left:56px; top:26px; width:1170px; font-size:38px; line-height:1.1;
            color:__TITLE2__; }
.slide ul { left:56px; top:160px; width:590px; }
.slide ul.wide { width:1170px; }
.slide ul.two { width:560px; }
.slide ul.right { left:650px; width:575px; }
.slide li { font-size:24px; line-height:1.34; margin:0 0 16px 28px; }
.slide li::before { content:"\\25B8"; left:-28px; }
.slide .pic { left:660px; top:160px; width:575px; height:470px; }
.slide .footname { left:56px; bottom:22px; }
.slide .badge { position:absolute; right:0; bottom:0; width:88px; height:52px;
                background:__ACCENT__; color:#ffffff; font-size:20px; font-weight:bold;
                display:flex; align-items:center; justify-content:center; }
.slide.title { background:linear-gradient(135deg, __TITLEFROM__, __TITLETO__); }
.slide.title h1 { left:96px; top:200px; width:1100px; color:__TITLE2__; }
.slide.title .sub { left:96px; top:330px; width:1000px; }
.slide.title .meta { left:96px; top:470px; width:1000px; }
""",
    'minimal': """
.slide .dot { position:absolute; left:56px; top:54px; width:16px; height:16px;
              background:__ACCENT__; }
.slide h1 { left:56px; top:88px; width:1020px; font-size:36px; line-height:1.2; }
.slide ul { left:56px; top:196px; width:590px; }
.slide ul.wide { width:1170px; }
.slide ul.two { width:560px; }
.slide ul.right { left:650px; width:575px; }
.slide li { font-size:21px; line-height:1.42; margin:0 0 16px 0; }
.slide li.dense { font-size:20px; line-height:1.38; margin:0 0 12px 0; }
.slide li::before { content:""; }
.slide .pic { left:660px; top:196px; width:575px; height:440px; }
.slide .pic img { border:1px solid #e2e8f0; }
.slide .bignum { position:absolute; right:56px; top:44px; font-size:76px; font-weight:bold;
                 color:#eef2f7; line-height:1; }
.slide .footname { left:56px; bottom:22px; }
.slide.title { background:__TITLEBG__; }
.slide.title .dot { left:96px; top:180px; }
.slide.title h1 { left:96px; top:215px; width:1100px; color:__TITLE2__; }
.slide.title .sub { left:96px; top:345px; width:1000px; }
.slide.title .meta { left:96px; top:480px; width:1000px; }
""",
}

JS = """
(function () {
  var slides = Array.prototype.slice.call(document.querySelectorAll('.slide'));
  var i = 0;
  var count = document.getElementById('count');
  function show(n) {
    i = Math.max(0, Math.min(slides.length - 1, n));
    slides.forEach(function (s, k) { s.classList.toggle('on', k === i); });
    count.textContent = (i + 1) + ' / ' + slides.length;
    if (location.hash !== '#' + (i + 1)) { history.replaceState(null, '', '#' + (i + 1)); }
  }
  function next() { show(i + 1); }
  function prev() { show(i - 1); }
  document.getElementById('next').addEventListener('click', next);
  document.getElementById('prev').addEventListener('click', prev);
  document.getElementById('first').addEventListener('click', function () { show(0); });
  document.getElementById('print').addEventListener('click', function () { window.print(); });
  document.addEventListener('keydown', function (e) {
    if (e.key === 'ArrowRight' || e.key === 'PageDown' || e.key === ' ' || e.key === 'Enter') {
      e.preventDefault(); next();
    } else if (e.key === 'ArrowLeft' || e.key === 'PageUp' || e.key === 'Backspace') {
      e.preventDefault(); prev();
    } else if (e.key === 'Home') { show(0); }
    else if (e.key === 'End') { show(slides.length - 1); }
  });
  document.getElementById('stage').addEventListener('click', function (e) {
    if (e.clientX > window.innerWidth * 0.25) { next(); } else { prev(); }
  });
  function fit() {
    var s = Math.min(window.innerWidth / 1280, window.innerHeight / 720);
    document.getElementById('scaler').style.transform = 'scale(' + s + ')';
  }
  window.addEventListener('resize', fit);
  fit();
  var start = parseInt((location.hash || '#1').slice(1), 10);
  show(isNaN(start) ? 0 : start - 1);
})();
"""


def css_for(theme):
    css = (BASE_CSS
           .replace('__BG__', '#' + theme['bg'])
           .replace('__FONT__', theme['font'])
           .replace('__TITLE__', '#' + theme['title_c'])
           .replace('__BODY__', '#' + theme['body_c'])
           .replace('__ACCENT__', '#' + theme['accent'])
           .replace('__MUTED__', '#' + theme['muted'])
           .replace('__RULE__', '#1e293b' if theme['decor'] == 'leftbar' else '#cbd5e1')
           .replace('__SUB__', '#' + theme['sub_c'])
           .replace('__META__', '#' + theme['meta_c'])
           .replace('__TITLEBG__', '#' + (theme['title_bg'] or '0b1020'))
           .replace('__TITLE2__', '#' + theme['title_c2']))
    extra = LAYOUT[theme['decor']]
    extra = (extra.replace('__ACCENT__', '#' + theme['accent'])
             .replace('__TITLEBG__', '#' + (theme['title_bg'] or '0b1020'))
             .replace('__TITLE2__', '#' + theme['title_c2'])
             .replace('__BAND1__', '#' + theme.get('band_from', theme['accent']))
             .replace('__BAND2__', '#' + theme.get('band_to', theme['accent']))
             .replace('__TITLEFROM__', '#' + theme.get('title_from', theme['accent']))
             .replace('__TITLETO__', '#' + theme.get('title_to', theme['accent'])))
    return css + extra


def slide_html(s, idx, theme, cache, deck_title):
    d = theme['decor']
    cls = 'slide title' if s['kind'] == 'title' else 'slide'
    if idx == 1:
        cls += ' on'   # первый слайд виден даже без JavaScript
    out = [f'<section class="{cls}" data-n="{idx}">']

    if s['kind'] == 'title':
        if d == 'frame':
            out.append('<div class="frame"></div>')
        if d in ('bar', 'leftbar'):
            out.append('<div class="stripe"></div>')
        if d == 'minimal':
            out.append('<div class="dot"></div>')
        out.append(f'<h1>{inline(s["title"])}</h1>')
        out.append(f'<div class="sub">{inline(s["sub"])}</div>')
        out.append('<div class="meta">' + '<br>'.join(inline(x) for x in s['meta']) + '</div>')
        out.append('</section>')
        return '\n'.join(out)

    if d == 'bar':
        out.append('<div class="topbar"></div>')
    elif d == 'leftbar':
        out.append('<div class="leftbar"></div>')
    elif d == 'frame':
        out.append('<div class="frame"></div>')
    elif d == 'band':
        out.append('<div class="band"></div>')
    elif d == 'minimal':
        out.append('<div class="dot"></div>')
        out.append(f'<div class="bignum">{idx}</div>')

    out.append(f'<h1>{inline(s["title"])}</h1>')
    if d in ('bar', 'leftbar'):
        out.append('<div class="rule"></div>')
    elif d == 'frame':
        out.append('<div class="rule"></div><div class="rule2"></div>')
    elif d == 'minimal':
        pass

    bullets = s.get('bullets', [])
    if s.get('img'):
        cls_li = ' class="dense"' if len(bullets) >= 6 else ''
        out.append('<ul>' + ''.join(f'<li{cls_li}>{inline(b)}</li>' for b in bullets) + '</ul>')
        src = 'data:image/png;base64,' + b64(s['img'], cache)
        out.append(f'<div class="pic"><img src="{src}" alt="{html.escape(s["title"])}"></div>')
    elif len(bullets) > 4:
        half = (len(bullets) + 1) // 2
        out.append('<ul class="two">' + ''.join(f'<li>{inline(b)}</li>'
                                                for b in bullets[:half]) + '</ul>')
        out.append('<ul class="right">' + ''.join(f'<li>{inline(b)}</li>'
                                                  for b in bullets[half:]) + '</ul>')
    else:
        out.append('<ul class="wide">' + ''.join(f'<li>{inline(b)}</li>'
                                                 for b in bullets) + '</ul>')

    if theme['number'] == 'footer':
        out.append('<div class="foot"></div>')
        out.append(f'<div class="footname">{deck_title}</div>')
        out.append(f'<div class="num">{idx}</div>')
    elif theme['number'] == 'badge':
        out.append(f'<div class="footname">{deck_title}</div>')
        out.append(f'<div class="badge">{idx}</div>')
    else:
        out.append(f'<div class="footname">{deck_title}</div>')

    out.append('</section>')
    return '\n'.join(out)


def build_html(theme, path, cache, deck):
    slides = deck_slides(deck)
    deck_title = html.escape(deck['title'])
    body = '\n'.join(slide_html(s, i, theme, cache, deck_title)
                      for i, s in enumerate(slides, start=1))
    text = (
        '<!DOCTYPE html>\n<html lang="ru">\n<head>\n<meta charset="utf-8">\n'
        f'<title>{deck_title} — {html.escape(deck["name"])}</title>\n'
        f'<style>{css_for(theme)}</style>\n</head>\n<body>\n'
        '<div id="stage"><div id="scaler">\n' + body + '\n</div></div>\n'
        '<div id="ui">'
        '<button id="first" title="На первый слайд">«</button>'
        '<button id="prev" title="Назад">←</button>'
        '<span id="count">1 / ' + str(len(slides)) + '</span>'
        '<button id="next" title="Вперёд">→</button>'
        '<button id="print" title="Печать">Печать</button>'
        '</div>\n'
        '<div id="hint">стрелки, пробел или щелчок — перелистывание; Home / End — начало и конец'
        '</div>\n'
        f'<script>{JS}</script>\n</body></html>\n'
    )
    with open(path, 'w', encoding='utf-8-sig') as fh:
        fh.write(text)
    return len(text)


def build_index(path):
    cards = []
    for i, deck in enumerate(DECKS, start=1):
        t = THEME_BY_KEY[deck['theme']]
        cards.append(
            f'  <li>\n'
            f'    <div class="n">{i}. {html.escape(deck["name"])}</div>\n'
            f'    <div class="who">{html.escape(deck["title"])} · '
            f'оформление: {html.escape(t["name"])}</div>\n'
            f'    <p>{html.escape(deck["desc"])}</p>\n'
            f'    <p class="links">\n'
            f'      <a href="{deck["file"]}.html">открыть в браузере</a> ·\n'
            f'      <a href="{deck["file"]}.pptx">скачать .pptx</a>\n'
            f'    </p>\n'
            f'  </li>')
    text = (
        '<!DOCTYPE html>\n<html lang="ru">\n<head>\n<meta charset="utf-8">\n'
        '<title>Пять презентаций о проекте Motion Playground Web</title>\n'
        '<style>\n'
        'body { margin:0; padding:40px 28px; background:#0b1020; color:#e5e7eb;\n'
        '       font-family:Verdana, sans-serif; line-height:1.5; }\n'
        'h1 { font-size:26px; margin:0 0 6px; }\n'
        '.sub { color:#93c5fd; margin:0 0 28px; }\n'
        'ul { list-style:none; padding:0; margin:0; display:grid; gap:16px;\n'
        '     grid-template-columns:repeat(auto-fit, minmax(340px, 1fr)); }\n'
        'li { background:#111827; border:1px solid #1f2937; border-radius:10px; padding:18px 20px; }\n'
        '.n { font-size:20px; font-weight:bold; color:#22d3ee; margin-bottom:6px; }\n'
        '.who { font-size:14px; color:#94a3b8; margin-bottom:10px; }\n'
        'p { margin:0 0 12px; font-size:15px; color:#cbd5e1; }\n'
        '.links { margin:0; }\n'
        'a { color:#93c5fd; }\n'
        '</style>\n</head>\n<body>\n'
        '<h1>Пять презентаций о проекте: разные аудитории, по 30 слайдов</h1>\n'
        '<p class="sub">Каждая презентация написана отдельно: свой текст, свои слайды '
        'и своё оформление.</p>\n'
        '<ul>\n' + '\n'.join(cards) + '\n</ul>\n'
        '</body></html>\n')
    with open(path, 'w', encoding='utf-8-sig') as fh:
        fh.write(text)
    return len(text)


# ---------------------------------------------------------------------------

def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    used = [s['img'] for deck in DECKS for s in deck['slides'] if s.get('img')]
    missing = sorted({n for n in used if not os.path.exists(os.path.join(IMG_DIR, n))})
    if missing:
        print('Нет рисунков:', ', '.join(missing))
        print('Сначала запустите: python3 tools/make_report_images.py')
        return 1
    cache = {}
    for i, deck in enumerate(DECKS, start=1):
        theme = THEME_BY_KEY[deck['theme']]
        slides = deck_slides(deck)
        h = os.path.join(OUT_DIR, deck['file'] + '.html')
        p = os.path.join(OUT_DIR, deck['file'] + '.pptx')
        hs = build_html(theme, h, cache, deck)
        print(f'  {i}. {deck["name"]}: {len(slides)} слайдов, '
              f'{sum(1 for x in deck["slides"] if x.get("img"))} рисунков')
        print(f'     {os.path.relpath(h, ROOT)}  ({hs // 1024} КБ)')
        if build_pptx(theme, p, slides, deck):
            print(f'     {os.path.relpath(p, ROOT)}  ({os.path.getsize(p) // 1024} КБ)')
    idx = os.path.join(OUT_DIR, 'index.html')
    build_index(idx)
    print(f'  {os.path.relpath(idx, ROOT)}  (оглавление)')
    print('Готово: 5 презентаций.')
    return 0


if __name__ == '__main__':
    sys.exit(main())
